"""Offline, lossless source normalization and candidate geography inventory.

No rent/value ratio, ownership scenario or geographic equivalence is inferred.
Run with the bundled Python runtime (openpyxl is read-only).
"""
from __future__ import annotations

import csv
import hashlib
import json
import re
import xml.etree.ElementTree as ET
from collections import Counter, defaultdict
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
NS = {"x": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}


def read_pipe(name):
    with (ROOT / "raw" / name).open(encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f, delimiter="|"))


def unique(rows, key):
    values = [row[key] for row in rows]
    if len(set(values)) != len(values):
        raise ValueError(f"Duplicate {key}; resolve source records before joining")


def numeric(raw):
    # Preserve original annotations/sentinel values in adjacent raw columns.
    try:
        value = int(raw)
    except (ValueError, TypeError):
        return "", "missing_or_annotation"
    if value < 0:
        return "", "negative_sentinel"
    return value, "numeric"


def write_csv(name, rows, fields):
    with (ROOT / "prepared" / name).open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)


def main():
    sources = json.loads((ROOT / "sources.json").read_text())
    for source in sources:
        raw = (ROOT / "raw" / source["file"]).read_bytes()
        if len(raw) != source["bytes"] or hashlib.sha256(raw).hexdigest() != source["sha256"]:
            raise ValueError(f"Source integrity changed: {source['file']}")

    geo = read_pipe("acs-2024-geographies.txt")
    values = read_pipe("acs-2024-b25077.dat")
    unique(geo, "GEO_ID")
    unique(values, "GEO_ID")
    estimates = {row["GEO_ID"]: row for row in values}
    cbsa_rows = [g for g in geo if g["SUMLEVEL"] == "310" and g["COMPONENT"] == "00"]
    cbsa = {g["CBSA"]: g for g in cbsa_rows}
    if len(cbsa) != len(cbsa_rows):
        raise ValueError("Duplicate CBSA in selected geography universe")

    delineation = load_workbook(ROOT / "raw" / "census-2023-cbsa-delineation.xlsx", read_only=True, data_only=True)
    expected_counties = defaultdict(set)
    try:
        for row in delineation["List 1"].iter_rows(min_row=4, values_only=True):
            if row[0] is not None and str(row[0]).isdigit():
                expected_counties[str(row[0]).zfill(5)].add(str(row[9]).zfill(2) + str(row[10]).zfill(3))
    finally:
        delineation.close()

    bookpath = ROOT / "raw" / "hud-fy2027-fmr.xlsx"
    # Formatting stretches the worksheet dimension to Excel's maximum row.
    # Inspect XML value-bearing cells; don't loop through a million empty rows.
    with ZipFile(bookpath) as z:
        tree = ET.fromstring(z.read("xl/worksheets/sheet1.xml"))
        last_row = max(int(r.attrib["r"]) for r in tree.findall(".//x:row", NS)
                       if any(c.find("x:v", NS) is not None or c.find("x:is", NS) is not None
                              for c in r.findall("x:c", NS)))
    wb = load_workbook(bookpath, read_only=True, data_only=True)
    try:
        sheet = wb["FY27_FMRs"]
        rows = list(sheet.iter_rows(min_row=1, max_row=last_row, max_col=14, values_only=True))
    finally:
        wb.close()
    fields = list(rows[0])
    required = ["stusps", "state", "hud_area_code", "countyname", "county_town_name", "metro",
                "hud_area_name", "fips", "pop2024", "fmr_0", "fmr_1", "fmr_2", "fmr_3", "fmr_4"]
    if fields != required:
        raise ValueError(f"Unexpected HUD schema: {fields}")
    hud = [dict(zip(fields, row)) for row in rows[1:] if any(v is not None for v in row)]
    groups = defaultdict(list)
    for row in hud:
        row["state"] = str(row["state"]).zfill(2)
        row["fips"] = str(row["fips"]).zfill(10)
        if not re.fullmatch(r"\d{10}", row["fips"]):
            raise ValueError("Unexpected HUD FIPS identifier")
        code = str(row["hud_area_code"])
        row["candidate_cbsa"] = code[5:10] if re.fullmatch(r"METRO\d{5}[MN].{5}", code) else ""
        row["source_fiscal_year"] = 2027
        groups[row["candidate_cbsa"]].append(row)

    metros = []
    for g in sorted(cbsa_rows, key=lambda row: row["CBSA"]):
        raw = estimates.get(g["GEO_ID"], {})
        estimate, est_status = numeric(raw.get("B25077_E001"))
        moe, moe_status = numeric(raw.get("B25077_M001"))
        members = groups.get(g["CBSA"], [])
        areas = sorted({r["hud_area_code"] for r in members})
        rents = sorted({r["fmr_2"] for r in members})
        expected = expected_counties.get(g["CBSA"], set())
        observed = {r["fips"][:5] for r in members}
        missing = sorted(expected - observed)
        extra = sorted(observed - expected)
        # County-set agreement alone cannot establish town/subcounty equivalence.
        coverage = "county_sets_equal_subcounty_unverified" if expected and not missing and not extra else "county_sets_differ_review"
        kind = "metro" if g["NAME"].endswith("Metro Area") else "micro" if g["NAME"].endswith("Micro Area") else "other"
        status = "no_hud_candidate" if not members else "multiple_hud_areas_review" if len(areas) > 1 else "single_hud_area_coverage_unverified"
        metros.append({"cbsa": g["CBSA"], "geo_id": g["GEO_ID"], "name": g["NAME"],
                       "area_kind": kind, "acs_year": 2024, "median_owner_value_raw": raw.get("B25077_E001", ""),
                       "median_owner_value_usd": estimate, "estimate_status": est_status,
                       "moe_raw": raw.get("B25077_M001", ""), "moe_usd": moe, "moe_status": moe_status,
                       "hud_candidate_records": len(members), "hud_area_count": len(areas),
                       "hud_area_codes": ";".join(areas), "hud_2br_values": ";".join(map(str, rents)),
                       "expected_county_count": len(expected), "observed_county_count": len(observed),
                       "missing_counties": ";".join(missing), "extra_counties": ";".join(extra),
                       "county_coverage_status": coverage,
                       "geography_status": status, "approved_for_model": "false"})
    orphan = [r for r in hud if r["candidate_cbsa"] and r["candidate_cbsa"] not in cbsa]
    out = ROOT / "prepared"
    out.mkdir(exist_ok=True)
    write_csv("hud-records.csv", hud, fields + ["candidate_cbsa", "source_fiscal_year"])
    write_csv("cbsa-candidates.csv", metros, list(metros[0]))
    write_csv("hud-unmatched-cbsa.csv", orphan, fields + ["candidate_cbsa", "source_fiscal_year"])
    report = {
        "status": "normalized candidates; no geographic equivalence or ownership results approved",
        "source_hashes_verified": len(sources), "hud_records": len(hud),
        "acs_geographies": len(geo), "acs_value_records": len(values),
        "geographies_without_value_row": sorted(set(g["GEO_ID"] for g in geo) - set(estimates)),
        "value_rows_without_geography": sorted(set(estimates) - set(g["GEO_ID"] for g in geo)),
        "cbsa_records": len(metros), "area_kinds": dict(Counter(r["area_kind"] for r in metros)),
        "metro_geography_status": dict(Counter(r["geography_status"] for r in metros if r["area_kind"] == "metro")),
        "metro_value_status": dict(Counter(r["estimate_status"] for r in metros if r["area_kind"] == "metro")),
        "metro_county_coverage": dict(Counter(r["county_coverage_status"] for r in metros if r["area_kind"] == "metro")),
        "hud_unmatched_candidate_records": len(orphan),
        "hud_unmatched_candidate_codes": sorted({r["candidate_cbsa"] for r in orphan}),
        "model_approved_rows": 0,
        "outputs": {p.name: {"bytes": p.stat().st_size, "sha256": hashlib.sha256(p.read_bytes()).hexdigest()}
                    for p in sorted(out.glob("*.csv"))},
    }
    (ROOT / "preparation.json").write_text(json.dumps(report, indent=2) + "\n")
    print(json.dumps(report, indent=2))


if __name__ == "__main__":
    main()
